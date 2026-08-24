"""PDF/DOCX -> plain text with page and section provenance preserved.

Design note: this deliberately does NOT special-case the synthetic demo
corpus. Section headings are recovered with simple, generic heuristics
(PDF: short line, no trailing punctuation; DOCX: paragraph style name) so
the same code path works on real regulatory PDFs/DOCX files dropped into
data/corpus/.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from pypdf import PdfReader

HEADING_MAX_CHARS = 90
_SENTENCE_END = re.compile(r"[.!?]\s*$")


@dataclass
class ParsedParagraph:
    text: str
    is_heading: bool
    page: int


@dataclass
class ParsedPage:
    page_number: int
    text: str


@dataclass
class ParsedDocument:
    doc_id: str
    title: str
    source_path: Path
    paragraphs: list[ParsedParagraph] = field(default_factory=list)
    pages: list[ParsedPage] = field(default_factory=list)


def _looks_like_heading(line: str) -> bool:
    line = line.strip()
    if not line or len(line) > HEADING_MAX_CHARS:
        return False
    if _SENTENCE_END.search(line):
        return False
    # Headings in this corpus are Title Case-ish and contain no more than
    # ~12 words; body sentences are longer and end in punctuation.
    return len(line.split()) <= 12


def parse_pdf(path: Path) -> ParsedDocument:
    reader = PdfReader(str(path))
    doc_id = path.stem
    paragraphs: list[ParsedParagraph] = []
    pages: list[ParsedPage] = []
    title = doc_id
    title_captured = False

    for page_idx, page in enumerate(reader.pages, start=1):
        raw = page.extract_text() or ""
        pages.append(ParsedPage(page_number=page_idx, text=raw))
        # pypdf does not reliably preserve blank lines between flowables, so
        # paragraphs are recovered by accumulating wrapped lines until one
        # ends in sentence-final punctuation (a heading, detected only at a
        # paragraph boundary, is emitted as its own single-line paragraph).
        lines = [ln.strip() for ln in raw.splitlines() if ln.strip()]
        buffer: list[str] = []
        blocks: list[str] = []
        mode: str | None = None  # "heading" | "body" | None
        for line in lines:
            heading_like = _looks_like_heading(line)
            if mode in (None, "heading") and heading_like:
                buffer.append(line)
                mode = "heading"
                continue
            if mode == "heading":
                # a wrapped multi-line heading ends where non-heading text begins
                blocks.append(" ".join(buffer))
                buffer = []
            buffer.append(line)
            mode = "body"
            if _SENTENCE_END.search(line):
                blocks.append(" ".join(buffer))
                buffer = []
                mode = None
        if buffer:
            blocks.append(" ".join(buffer))

        for block in blocks:
            normalized = " ".join(block.split())
            if not normalized:
                continue
            if page_idx == 1 and not title_captured:
                title = normalized
                title_captured = True
                continue
            is_heading = _looks_like_heading(normalized)
            paragraphs.append(ParsedParagraph(text=normalized, is_heading=is_heading, page=page_idx))

    return ParsedDocument(doc_id=doc_id, title=title, source_path=path, paragraphs=paragraphs, pages=pages)


def parse_docx(path: Path) -> ParsedDocument:
    document = DocxDocument(str(path))
    doc_id = path.stem
    paragraphs: list[ParsedParagraph] = []
    page_texts: dict[int, list[str]] = {1: []}
    page = 1
    title = doc_id
    seen_title = False

    for para in document.paragraphs:
        # A run-level <w:br w:type="page"/> is an explicit page break.
        for run in para.runs:
            breaks = run._element.findall(qn("w:br"))
            for br in breaks:
                if br.get(qn("w:type")) == "page":
                    page += 1
                    page_texts.setdefault(page, [])

        text = para.text.strip()
        if not text:
            continue

        style_name = (para.style.name or "").lower() if para.style else ""
        is_heading = style_name.startswith("heading") or style_name == "title"

        if style_name == "title" and not seen_title:
            title = text
            seen_title = True
            page_texts[page].append(text)
            continue

        paragraphs.append(ParsedParagraph(text=text, is_heading=is_heading, page=page))
        page_texts.setdefault(page, []).append(text)

    pages = [ParsedPage(page_number=p, text="\n".join(t)) for p, t in sorted(page_texts.items())]
    return ParsedDocument(doc_id=doc_id, title=title, source_path=path, paragraphs=paragraphs, pages=pages)


def parse_document(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    raise ValueError(f"Unsupported document type: {suffix} ({path})")
